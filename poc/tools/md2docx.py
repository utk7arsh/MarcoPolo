"""Markdown (the subset our docs use) to a .docx with real Word styles, so Google Docs imports it cleanly.

- Headings become Title / Heading 1 / Heading 2; the paragraph right after the title is a muted subtitle.
- Lists use real numbering definitions: nesting follows the indent, every list (and every sub-list
  under a new parent item) restarts at 1, and bullets and numbers can be mixed across levels.
- Tables get borders and a shaded header row; fenced code goes in a shaded one-cell table, monospace.
- Inline **bold**, *italic* and `code`.
Usage: md2docx.py <in.md> <out.docx>
"""
import re
import sys

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

FONT, MONO = "Arial", "Courier New"
INK, ACCENT, MUTED = RGBColor(0x1F, 0x1F, 0x1F), RGBColor(0x1A, 0x4F, 0x8B), RGBColor(0x5F, 0x63, 0x68)
HEADER_FILL, CODE_FILL = "E8EEF7", "F3F4F6"
LIST_RE = re.compile(r"^( *)([-*]|\d+\.) (.*)$")
INLINE_RE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|\*[^*\s][^*]*\*)")
INDENT_TWIPS, HANG_TWIPS = 360, 300
BULLETS = ["•", "◦", "▪"]


def set_font(el_rpr_owner, name, size=None, color=None, bold=None, italic=None):
    """Font on a style or run, with the theme-font attributes removed (they would override the name)."""
    font = el_rpr_owner.font
    font.name = name
    rpr = el_rpr_owner.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        rfonts.attrib.pop(qn(attr), None)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(attr), name)
    if size is not None:
        font.size = Pt(size)
    if color is not None:
        font.color.rgb = color
    if bold is not None:
        font.bold = bold
    if italic is not None:
        font.italic = italic


def setup_styles(doc):
    st = doc.styles
    set_font(st["Normal"], FONT, 11, INK)
    pf = st["Normal"].paragraph_format
    pf.space_after, pf.line_spacing = Pt(6), 1.15
    for name, size, color, before, after in (("Title", 22, INK, 0, 4), ("Heading 1", 16, ACCENT, 18, 6), ("Heading 2", 13, INK, 12, 4)):
        s = st[name]
        set_font(s, FONT, size, color, bold=True, italic=False)
        s.paragraph_format.space_before, s.paragraph_format.space_after = Pt(before), Pt(after)
        s.paragraph_format.keep_with_next = True
        ppr = s.element.get_or_add_pPr()
        for b in ppr.findall(qn("w:pBdr")):  # the default Title has a bottom border; drop it
            ppr.remove(b)
    lp = st["List Paragraph"]
    set_font(lp, FONT, 11, INK)
    lp.paragraph_format.space_after = Pt(2)
    sec = doc.sections[0]
    sec.left_margin = sec.right_margin = Inches(0.9)
    sec.top_margin = sec.bottom_margin = Inches(0.85)


class Numbering:
    """Two abstract lists (all-bullet, all-decimal) and a fresh <w:num> per list run, so runs restart at 1."""

    def __init__(self, doc):
        self.root = doc.part.numbering_part.element
        self.next_abs, self.next_num = 900, 900
        self.abstract = {kind: self._abstract(kind) for kind in ("bullet", "decimal")}

    def _abstract(self, kind):
        a = OxmlElement("w:abstractNum")
        a.set(qn("w:abstractNumId"), str(self.next_abs))
        self.next_abs += 1
        mlt = OxmlElement("w:multiLevelType")
        mlt.set(qn("w:val"), "multilevel")
        a.append(mlt)
        for lvl in range(9):
            el = OxmlElement("w:lvl")
            el.set(qn("w:ilvl"), str(lvl))
            for tag, val in (("w:start", "1"), ("w:numFmt", "bullet" if kind == "bullet" else "decimal"),
                             ("w:lvlText", BULLETS[lvl % 3] if kind == "bullet" else f"%{lvl + 1}."), ("w:lvlJc", "left")):
                e = OxmlElement(tag)
                e.set(qn("w:val"), val)
                el.append(e)
            ppr = OxmlElement("w:pPr")
            ind = OxmlElement("w:ind")
            ind.set(qn("w:left"), str(INDENT_TWIPS * (lvl + 1)))
            ind.set(qn("w:hanging"), str(HANG_TWIPS))
            ppr.append(ind)
            el.append(ppr)
            rpr = OxmlElement("w:rPr")
            rf = OxmlElement("w:rFonts")
            for attr in ("w:ascii", "w:hAnsi", "w:cs"):
                rf.set(qn(attr), FONT)
            rpr.append(rf)
            el.append(rpr)
            a.append(el)
        first_num = self.root.find(qn("w:num"))
        if first_num is not None:
            first_num.addprevious(a)
        else:
            self.root.append(a)
        return self.next_abs - 1

    def new_list(self, kind, level):
        n = OxmlElement("w:num")
        n.set(qn("w:numId"), str(self.next_num))
        self.next_num += 1
        ab = OxmlElement("w:abstractNumId")
        ab.set(qn("w:val"), str(self.abstract[kind]))
        n.append(ab)
        ov = OxmlElement("w:lvlOverride")
        ov.set(qn("w:ilvl"), str(level))
        so = OxmlElement("w:startOverride")
        so.set(qn("w:val"), "1")
        ov.append(so)
        n.append(ov)
        self.root.append(n)
        return self.next_num - 1


def add_inline(par, text, size=None):
    for tok in INLINE_RE.split(text):
        if not tok:
            continue
        if tok.startswith("**") and tok.endswith("**"):
            run = par.add_run(tok[2:-2]); run.bold = True
        elif tok.startswith("`") and tok.endswith("`"):
            run = par.add_run(tok[1:-1]); set_font(run, MONO, 9.5, RGBColor(0x37, 0x41, 0x51))
            continue
        elif tok.startswith("*") and tok.endswith("*") and len(tok) > 2:
            run = par.add_run(tok[1:-1]); run.italic = True
        else:
            run = par.add_run(tok)
        if size:
            run.font.size = Pt(size)


def set_numbering(par, num_id, level):
    ppr = par._p.get_or_add_pPr()
    numpr = OxmlElement("w:numPr")
    il = OxmlElement("w:ilvl"); il.set(qn("w:val"), str(level))
    ni = OxmlElement("w:numId"); ni.set(qn("w:val"), str(num_id))
    numpr.append(il); numpr.append(ni)
    ppr.insert(0, numpr)


def shade(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), fill)
    tcpr.append(shd)


def add_table(doc, rows):
    header, body = rows[0], rows[1:]
    ncol = len(header)
    t = doc.add_table(rows=1 + len(body), cols=ncol)
    t.style = doc.styles["Table Grid"]
    t.autofit = False
    lens = [max(len(r[c]) if c < len(r) else 0 for r in rows) for c in range(ncol)]
    widths = [max(0.9, 6.7 * min(l, 90) / max(1, sum(min(x, 90) for x in lens))) for l in lens]
    scale = 6.7 / sum(widths)
    for ri, row in enumerate([header] + body):
        for ci in range(ncol):
            cell = t.cell(ri, ci)
            cell.width = Inches(widths[ci] * scale)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            add_inline(p, row[ci] if ci < len(row) else "", size=10)
            if ri == 0:
                for r in p.runs:
                    r.bold = True
                shade(cell, HEADER_FILL)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_code(doc, lines):
    t = doc.add_table(rows=1, cols=1)
    t.autofit = False
    cell = t.cell(0, 0)
    cell.width = Inches(6.7)
    shade(cell, CODE_FILL)
    first = True
    for line in lines:
        p = cell.paragraphs[0] if first else cell.add_paragraph()
        first = False
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        run = p.add_run(line if line else " ")
        set_font(run, MONO, 8.5, RGBColor(0x1F, 0x29, 0x37))
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def convert(src, dst):
    doc = Document()
    setup_styles(doc)
    numbering = Numbering(doc)
    lines = open(src, encoding="utf-8").read().splitlines()
    runs = {}            # level -> (kind, numId) of the list run currently open at that level
    after_title = False
    i = 0
    while i < len(lines):
        line = lines[i]
        if line.startswith("```"):
            j = i + 1
            while j < len(lines) and not lines[j].startswith("```"):
                j += 1
            add_code(doc, lines[i + 1:j])
            runs.clear(); i = j + 1; continue
        if line.startswith("|") and i + 1 < len(lines) and re.match(r"^\|[\s:|-]+\|$", lines[i + 1].strip()):
            rows, j = [], i
            while j < len(lines) and lines[j].startswith("|"):
                if j != i + 1:
                    rows.append([c.strip() for c in lines[j].strip().strip("|").split("|")])
                j += 1
            add_table(doc, rows)
            runs.clear(); i = j; continue
        m = LIST_RE.match(line)
        if m:
            n = len(m.group(1))
            level = 0 if n == 0 else 1 if n <= 3 else 2 if n <= 6 else 3
            kind = "decimal" if m.group(2)[0].isdigit() else "bullet"
            for lv in [k for k in runs if k > level]:
                del runs[lv]       # a new parent item: its sub-lists restart
            if level not in runs or runs[level][0] != kind:
                runs[level] = (kind, numbering.new_list(kind, level))
            p = doc.add_paragraph(style="List Paragraph")
            set_numbering(p, runs[level][1], level)
            add_inline(p, m.group(3))
            i += 1; continue
        stripped = line.strip()
        runs.clear()
        if not stripped or stripped == "---":
            i += 1; continue
        if line.startswith("# "):
            doc.add_paragraph(line[2:].strip(), style="Title"); after_title = True
        elif line.startswith("## "):
            doc.add_paragraph(line[3:].strip(), style="Heading 1")
        elif line.startswith("### "):
            doc.add_paragraph(line[4:].strip(), style="Heading 2")
        else:
            p = doc.add_paragraph()
            add_inline(p, stripped)
            if after_title:
                for r in p.runs:
                    r.font.color.rgb = MUTED; r.italic = True
                p.paragraph_format.space_after = Pt(12)
            after_title = False
        i += 1
    doc.save(dst)


if __name__ == "__main__":
    convert(sys.argv[1], sys.argv[2])
    print(sys.argv[2])
